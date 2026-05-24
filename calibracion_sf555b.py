"""
Calibración SF555B — toma muestras cada 2 s y exporta a Excel
para comparar con el sonómetro y calcular el ajuste.
"""

import sounddevice as sd
import numpy as np
import time
import sys
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SAMPLE_RATE = 44100
BLOCK_SIZE  = 1024
INTERVALO   = 2  # segundos entre muestras


def listar_micros():
    print("\n=== Dispositivos de entrada disponibles ===")
    encontrados = []
    for i, d in enumerate(sd.query_devices()):
        if d["max_input_channels"] > 0:
            print(f"  ID {i:2d} | canales: {d['max_input_channels']} | {d['name']}")
            encontrados.append(i)
    if not encontrados:
        print("  [!] No se encontraron dispositivos.")
    return encontrados


def abrir_micro(device_id):
    info = sd.query_devices(device_id)
    if int(info["max_input_channels"]) == 0:
        raise ValueError(f"ID {device_id} no tiene canales de entrada.")
    print(f"[OK] '{info['name']}' — 1 canal a {SAMPLE_RATE} Hz\n")
    return sd.InputStream(
        device=device_id,
        channels=1,
        samplerate=SAMPLE_RATE,
        blocksize=BLOCK_SIZE,
        dtype="float32",
    )


def medir_db(stream, segundos=2):
    n_bloques = max(1, int(SAMPLE_RATE * segundos / BLOCK_SIZE))
    cuadrados = []
    for _ in range(n_bloques):
        bloque, _ = stream.read(BLOCK_SIZE)
        cuadrados.append(np.mean(bloque[:, 0] ** 2))
    rms = np.sqrt(np.mean(cuadrados))
    return 20.0 * np.log10(rms + 1e-9)


def exportar_xlsx(muestras, archivo):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calibración"

    # Encabezados
    encabezados = ["#", "Hora", "Mic (dBFS)", "Sonómetro (dB)", "Diferencia (dB)"]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    for col, texto in enumerate(encabezados, 1):
        c = ws.cell(row=1, column=col, value=texto)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    # Datos
    for fila, (ts, db_mic) in enumerate(muestras, 2):
        ws.cell(row=fila, column=1, value=fila - 1)
        ws.cell(row=fila, column=2, value=ts)
        ws.cell(row=fila, column=3, value=round(db_mic, 2))
        ws.cell(row=fila, column=4, value=None)   # el usuario llena el sonómetro
        # Fórmula: Sonómetro - Mic
        ws.cell(row=fila, column=5, value=f"=D{fila}-C{fila}")

    # Fila de promedio
    n = len(muestras) + 1
    fila_prom = n + 2
    ws.cell(row=fila_prom, column=2, value="PROMEDIO OFFSET").font = Font(bold=True)
    ws.cell(row=fila_prom, column=5, value=f"=AVERAGE(E2:E{n})").font = Font(bold=True)

    # Ancho de columnas
    for col, ancho in zip("ABCDE", [6, 12, 16, 18, 18]):
        ws.column_dimensions[col].width = ancho

    wb.save(archivo)
    print(f"\n[OK] Guardado: {archivo}")
    print(f"     Rellena la columna 'Sonómetro (dB)' con tus lecturas.")
    print(f"     La columna 'Diferencia' y el promedio se calculan solos.")


def main():
    ids = listar_micros()
    if not ids:
        sys.exit(1)

    try:
        mic_id = int(input("\nIngrese el ID del micrófono USB: "))
    except ValueError:
        sys.exit("[!] ID inválido.")
    if mic_id not in ids:
        sys.exit(f"[!] ID {mic_id} no válido.")

    try:
        stream = abrir_micro(mic_id)
    except ValueError as e:
        sys.exit(f"[ERROR] {e}")

    print("Midiendo cada 2 segundos — presione Ctrl+C para detener y guardar.\n")
    print(f"{'#':>4}  {'Hora':>10}  {'dBFS':>8}")
    print("-" * 28)

    muestras = []
    try:
        stream.start()
        n = 1
        while True:
            db = medir_db(stream, segundos=INTERVALO)
            hora = datetime.now().strftime("%H:%M:%S")
            muestras.append((hora, db))
            print(f"{n:>4}  {hora:>10}  {db:8.2f}")
            n += 1
    except KeyboardInterrupt:
        print(f"\n--- {len(muestras)} muestras capturadas ---")
    finally:
        stream.stop()
        stream.close()

    if muestras:
        nombre = f"calibracion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        exportar_xlsx(muestras, nombre)


if __name__ == "__main__":
    main()
